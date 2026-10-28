
# -*- coding: utf-8 -*-
"""
DVIZH Material Downloader — первая версия.

Структура:
Категория / Номер заявки / Материал / скачанные файлы

Запуск: двойной щелчок по файлу .py.
Исходный Excel не изменяется.
"""

from pathlib import Path, PurePosixPath
from urllib.parse import urlparse, unquote, parse_qsl
from datetime import datetime
import csv
import mimetypes
import re
import time
import traceback
import tkinter as tk
from tkinter import filedialog, messagebox
import shutil
import subprocess
import sys

try:
    import requests
    from openpyxl import load_workbook
except ImportError as exc:
    print("Не найдена библиотека:", exc)
    print("Нужны requests и openpyxl.")
    input("\nНажми Enter для закрытия...")
    raise SystemExit(1)

try:
    import gdown
except ImportError:
    gdown = None


APP_NAME = "DVIZH Material Downloader"
MAX_RETRIES = 3
CHUNK_SIZE = 256 * 1024
TIMEOUT = (20, 90)

# YouTube: прокси, с которым была проверена загрузка через yt-dlp.
YOUTUBE_PROXY = "http://127.0.0.1:10809"

USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 Chrome/124.0 Safari/537.36"
)

URL_RE = re.compile(r'https?://[^\s<>"\'\]\[(){}]+', re.IGNORECASE)


def normalize(value):
    return re.sub(
        r"\s+", " ",
        str(value or "").strip().lower().replace("ё", "е")
    )


def safe_name(value, fallback="Без названия"):
    name = str(value or "").strip()
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name)
    name = re.sub(r"\s+", " ", name).strip(" .")
    return (name[:100].rstrip(" .") or fallback)


def unique_path(path):
    """Не перезаписывать существующие файлы."""
    path = Path(path)
    if not path.exists():
        return path

    index = 2
    while True:
        candidate = path.with_name(
            f"{path.stem} ({index}){path.suffix}"
        )
        if not candidate.exists():
            return candidate
        index += 1


def choose_file():
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    try:
        filename = filedialog.askopenfilename(
            title="Выбери Excel-таблицу со ссылками",
            filetypes=[
                ("Excel", "*.xlsx *.xlsm"),
                ("Все файлы", "*.*")
            ]
        )
    finally:
        root.destroy()

    return Path(filename) if filename else None


def choose_folder():
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    try:
        folder = filedialog.askdirectory(
            title="Выбери папку для скачивания"
        )
    finally:
        root.destroy()

    return Path(folder) if folder else None


def find_column(headers, candidates):
    headers_norm = [normalize(h) for h in headers]

    for candidate in candidates:
        candidate = normalize(candidate)
        for index, header in enumerate(headers_norm):
            if header == candidate:
                return index

    for candidate in candidates:
        candidate = normalize(candidate)
        for index, header in enumerate(headers_norm):
            if candidate and candidate in header:
                return index

    return None


def extract_urls(value, hyperlink=None):
    urls = []

    if hyperlink and hyperlink.lower().startswith(("http://", "https://")):
        urls.append(hyperlink.strip())

    if value is not None:
        for url in URL_RE.findall(str(value)):
            url = url.rstrip(".,;:!?")
            if url and url not in urls:
                urls.append(url)

    return urls


def read_excel(excel_path):
    """
    Ищет заголовки в первых десяти строках каждого листа.
    Нужен столбец Ссылка, URL или Link.
    """
    workbook = load_workbook(excel_path, data_only=True)
    records = []

    try:
        for sheet in workbook.worksheets:
            header_row = None
            url_col = None

            for row_num in range(1, min(sheet.max_row, 10) + 1):
                values = [
                    sheet.cell(row_num, col).value
                    for col in range(1, sheet.max_column + 1)
                ]
                candidate = find_column(
                    values, ["Ссылка", "URL", "Link"]
                )
                if candidate is not None:
                    header_row = row_num
                    url_col = candidate
                    break

            if header_row is None:
                continue

            headers = [
                sheet.cell(header_row, col).value
                for col in range(1, sheet.max_column + 1)
            ]

            no_col = find_column(
                headers,
                ["№ заявки", "Номер заявки", "Номер выступления"]
            )
            id_col = find_column(
                headers, ["ID заявки", "Request ID"]
            )
            category_col = find_column(
                headers,
                ["Категория", "Номинация", "Тема заявки", "Раздел"]
            )
            material_col = find_column(
                headers,
                ["Материал", "Тип материала", "Название материала",
                 "Исходное поле"]
            )

            for row_num in range(header_row + 1, sheet.max_row + 1):
                cell = sheet.cell(row_num, url_col + 1)

                urls = extract_urls(
                    cell.value,
                    cell.hyperlink.target if cell.hyperlink else None
                )
                if not urls:
                    continue

                def cell_value(column):
                    if column is None:
                        return None
                    return sheet.cell(row_num, column + 1).value

                request_no = cell_value(no_col)
                if request_no is None or not str(request_no).strip():
                    request_no = cell_value(id_col)

                category = cell_value(category_col)
                material = cell_value(material_col)

                record = {
                    "request_no": str(request_no or "Без номера").strip(),
                    "category": str(category or "Без категории").strip(),
                    "material": str(material or "Материал").strip(),
                    "sheet": sheet.title,
                    "row": row_num,
                }

                for url in urls:
                    records.append({**record, "url": url})

    finally:
        workbook.close()

    # Убираем точные дубликаты ссылок в одинаковом контексте.
    result = []
    seen = set()

    for item in records:
        key = (
            item["request_no"], item["category"],
            item["material"], item["url"]
        )
        if key not in seen:
            seen.add(key)
            result.append(item)

    return result


def log_result(log_path, item, status, saved_path="",
               size="", attempts="", details=""):
    fields = [
        "time", "status", "category", "request_no",
        "material", "url", "saved_path", "bytes",
        "attempts", "details"
    ]

    exists = log_path.exists()

    row = {
        "time": datetime.now().isoformat(timespec="seconds"),
        "status": status,
        "category": item.get("category", ""),
        "request_no": item.get("request_no", ""),
        "material": item.get("material", ""),
        "url": item.get("url", ""),
        "saved_path": str(saved_path),
        "bytes": size,
        "attempts": attempts,
        "details": details,
    }

    with log_path.open("a", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=fields)
        if not exists:
            writer.writeheader()
        writer.writerow(row)

    print(
        f"[{status}] Заявка {row['request_no']} | "
        f"{row['material']} | {details or saved_path}"
    )


def filename_from_response(response, original_url, fallback):
    disposition = response.headers.get("Content-Disposition", "")
    match = re.search(
        r"filename\*=UTF-8''([^;]+)", disposition, re.IGNORECASE
    )

    if match:
        name = unquote(match.group(1))
    else:
        match = re.search(
            r'filename="?([^";]+)', disposition, re.IGNORECASE
        )
        name = unquote(match.group(1)) if match else ""

    if not name:
        name = unquote(
            PurePosixPath(
                urlparse(response.url or original_url).path
            ).name
        )

    name = safe_name(name, fallback)

    # Добавляем расширение только если сервер сообщил подходящий тип.
    if "." not in Path(name).name:
        content_type = response.headers.get(
            "Content-Type", ""
        ).split(";")[0].strip()
        extension = mimetypes.guess_extension(content_type)
        if extension and extension not in (".html", ".htm"):
            name += extension

    return name


def download_http(url, folder, session, item, log_path):
    """Потоковая загрузка с повторами и временным файлом."""
    last_error = ""

    for attempt in range(1, MAX_RETRIES + 1):
        temp_path = None

        try:
            with session.get(
                url, stream=True, timeout=TIMEOUT,
                allow_redirects=True
            ) as response:
                response.raise_for_status()

                content_type = response.headers.get(
                    "Content-Type", ""
                ).lower()

                if "text/html" in content_type:
                    raise ValueError(
                        "Получена HTML-страница вместо файла. "
                        "Ссылка может требовать доступа."
                    )

                filename = filename_from_response(
                    response, url, item["material"]
                )
                target = unique_path(folder / filename)
                temp_path = target.with_name(target.name + ".part")

                length = response.headers.get("Content-Length")
                expected = int(length) if length and length.isdigit() else None
                total = 0

                with temp_path.open("wb") as output:
                    for chunk in response.iter_content(
                        chunk_size=CHUNK_SIZE
                    ):
                        if chunk:
                            output.write(chunk)
                            total += len(chunk)

                if expected is not None and total != expected:
                    raise IOError(
                        f"Размер не совпал: ожидалось {expected}, "
                        f"получено {total} байт"
                    )

                temp_path.replace(target)

                log_result(
                    log_path, item, "OK", target, total,
                    attempt, "Скачано"
                )
                return [target]

        except Exception as exc:
            last_error = f"{type(exc).__name__}: {exc}"

            if temp_path and temp_path.exists():
                try:
                    temp_path.unlink()
                except OSError:
                    pass

            print(f"  Попытка {attempt}/{MAX_RETRIES}: {last_error}")

            if attempt < MAX_RETRIES:
                time.sleep(2 * attempt)

    log_result(
        log_path, item, "ERROR",
        attempts=MAX_RETRIES, details=last_error
    )
    return []


# ============================================================
# YouTube
# ============================================================

def is_youtube(url):
    """Распознаёт ссылки youtube.com, youtu.be и youtube-nocookie.com."""
    host = (urlparse(url).hostname or "").lower()

    return (
        host == "youtu.be"
        or host == "youtube.com"
        or host.endswith(".youtube.com")
        or host == "youtube-nocookie.com"
        or host.endswith(".youtube-nocookie.com")
    )


def download_youtube(url, folder, item, log_path):
    """Надёжная загрузка YouTube через yt-dlp с повторными попытками."""

    folder.mkdir(parents=True, exist_ok=True)

    max_attempts = 3
    temp_folder = folder / (
        f".youtube_tmp_{int(time.time() * 1000)}"
    )
    temp_folder.mkdir(parents=True, exist_ok=True)

    output_template = str(
        temp_folder / "%(title)s.%(ext)s"
    )

    last_error = None

    try:
        for attempt in range(1, max_attempts + 1):
            print(
                f"\n[YouTube] Попытка {attempt}/{max_attempts}"
            )
            print(f"[YouTube] Ссылка: {url}")

            command = [
                sys.executable,
                "-m", "yt_dlp",

                "--proxy", YOUTUBE_PROXY,
                "--no-playlist",

                # Клиенты, с которыми прошёл тест.
                "--extractor-args",
                "youtube:player_client=visionos,web_embedded,mweb",

                # Повторные попытки при временных сбоях.
                "--retries", "10",
                "--fragment-retries", "10",
                "--file-access-retries", "3",
                "--socket-timeout", "30",

                # Сначала MP4, иначе лучший доступный формат.
                "-f", "best[ext=mp4]/best",

                # Временные файлы остаются внутри временной папки.
                "-o", output_template,

                url,
            ]

            try:
                result = subprocess.run(
                    command,
                    check=False,
                )

                if result.returncode != 0:
                    raise RuntimeError(
                        f"yt-dlp завершился с кодом "
                        f"{result.returncode}"
                    )

                # Ищем готовый файл, исключая незавершённые загрузки.
                downloaded_files = [
                    p for p in temp_folder.iterdir()
                    if p.is_file()
                    and not p.name.endswith(
                        (".part", ".ytdl", ".temp")
                    )
                ]

                if not downloaded_files:
                    raise RuntimeError(
                        "yt-dlp завершился без ошибки, "
                        "но готовый файл не найден."
                    )

                # При обычной загрузке ожидается один файл.
                # Если yt-dlp создал несколько, переносим все.
                saved_files = []

                for source in downloaded_files:
                    if source.stat().st_size <= 0:
                        raise RuntimeError(
                            f"Получен пустой файл: {source.name}"
                        )

                    target = unique_path(folder / source.name)
                    source.replace(target)
                    saved_files.append(target)

                total_size = sum(
                    p.stat().st_size for p in saved_files
                )

                log_result(
                    log_path,
                    item,
                    "OK",
                    saved_files[0],
                    total_size,
                    attempt,
                    "YouTube: загрузка завершена"
                )

                print(
                    f"[YouTube] Готово: {saved_files[0].name}"
                )

                return saved_files

            except Exception as exc:
                last_error = exc
                print(
                    f"[YouTube] Ошибка попытки "
                    f"{attempt}: {exc}"
                )

                if attempt < max_attempts:
                    print(
                        "[YouTube] Повтор через 3 секунды..."
                    )
                    time.sleep(3)

        raise RuntimeError(
            f"YouTube: не удалось скачать после "
            f"{max_attempts} попыток. Последняя ошибка: "
            f"{last_error}"
        )

    finally:
        # Удаляем временную папку и недокачанные .part-файлы.
        shutil.rmtree(temp_folder, ignore_errors=True)


# ============================================================
# Яндекс Диск и Google Drive
# ============================================================

def is_yandex(url):
    host = urlparse(url).netloc.lower()
    return (
        host.endswith("yandex.ru")
        or host.endswith("yadi.sk")
        or host.endswith("disk.yandex.com")
    )


def is_google_drive(url):
    host = urlparse(url).netloc.lower()
    return "drive.google.com" in host or "docs.google.com" in host



def yandex_download_url(public_url, session):
    api_url = (
        "https://cloud-api.yandex.net/v1/disk/public/resources/download"
    )
    last_error = None

    for attempt in range(1, MAX_RETRIES + 1):
        try:
            response = session.get(
                api_url,
                params={"public_key": public_url},
                timeout=TIMEOUT
            )
            response.raise_for_status()
            data = response.json()

            if not data.get("href"):
                raise RuntimeError(
                    "Яндекс Диск не вернул ссылку скачивания."
                )

            return data["href"], data

        except requests.HTTPError as exc:
            last_error = exc
            status = (
                exc.response.status_code
                if exc.response is not None else None
            )

            # Повторяем только временные ошибки сервера.
            if status is None or status < 500 or attempt == MAX_RETRIES:
                raise

        except (requests.Timeout, requests.ConnectionError) as exc:
            last_error = exc
            if attempt == MAX_RETRIES:
                raise

        if attempt < MAX_RETRIES:
            delay = 2 * attempt
            print(
                f"  Яндекс Диск: повтор {attempt + 1}/"
                f"{MAX_RETRIES} через {delay} сек."
            )
            time.sleep(delay)

    raise RuntimeError(
        f"Не удалось получить ссылку Яндекс Диска: {last_error}"
    )



def google_drive_id(url):
    parsed = urlparse(url)

    match = re.search(r"/folders/([A-Za-z0-9_-]+)", parsed.path)
    if match:
        return match.group(1), True

    match = re.search(r"/file/d/([A-Za-z0-9_-]+)", parsed.path)
    if match:
        return match.group(1), False

    query = dict(parse_qsl(parsed.query))
    if query.get("id"):
        return query["id"], False

    return None, False

def download_google(url, folder, item, log_path):
    if gdown is None:
        raise RuntimeError(
            "Не найдена библиотека gdown.\n"
            "Установи её командой:\n"
            "python -m pip install -U gdown"
        )

    file_id, is_folder = google_drive_id(url)

    if not file_id:
        raise ValueError(
            "Не удалось определить ID Google Drive."
        )

    folder.mkdir(parents=True, exist_ok=True)

    temp_folder = folder / (
        f".gdrive_tmp_{int(time.time() * 1000)}"
    )
    temp_folder.mkdir(parents=True, exist_ok=True)

    try:
        if is_folder:
            result = gdown.download_folder(
                url=url,
                output=str(temp_folder),
                quiet=False,
                use_cookies=False,
            )

            # Ищем все полученные файлы внутри временной папки.
            downloaded = [
                p for p in temp_folder.rglob("*")
                if p.is_file()
                and p.stat().st_size > 0
                and not p.name.endswith(
                    (".part", ".aria2", ".tmp")
                )
            ]

            if not downloaded:
                raise RuntimeError(
                    "gdown не обнаружил скачанных файлов. "
                    "Возможно, папка закрыта или недоступна."
                )

            saved_files = []

            for source in downloaded:
                relative = source.relative_to(temp_folder)
                destination_dir = folder / relative.parent
                destination_dir.mkdir(
                    parents=True, exist_ok=True
                )

                target = unique_path(
                    destination_dir / relative.name
                )
                shutil.move(str(source), str(target))
                saved_files.append(target)

            total_size = sum(
                p.stat().st_size for p in saved_files
            )

            log_result(
                log_path,
                item,
                "OK",
                folder,
                total_size,
                1,
                f"Google Drive: скачано файлов: "
                f"{len(saved_files)}",
            )

            return saved_files

        # Одиночный файл: сохраняем во временную папку,
        # чтобы не записать HTML-страницу под именем материала.
        requested_name = safe_name(
            item["material"], "GoogleDrive_file"
        )
        temp_file = temp_folder / requested_name

        result = gdown.download(
            url=url,
            output=str(temp_file),
            quiet=False,
            use_cookies=False,
            resume=True,
        )

        if not result:
            raise RuntimeError(
                "gdown не смог скачать файл Google Drive."
            )

        result_path = Path(result)

        if not result_path.is_absolute():
            result_path = temp_folder / result_path

        if not result_path.exists():
            # Некоторые версии gdown могут сохранить файл
            # под именем, полученным от Google.
            candidates = [
                p for p in temp_folder.rglob("*")
                if p.is_file()
                and p.stat().st_size > 0
            ]
            if len(candidates) == 1:
                result_path = candidates[0]
            else:
                raise RuntimeError(
                    "gdown сообщил об успехе, но файл "
                    "не найден."
                )

        if result_path.stat().st_size <= 0:
            raise RuntimeError(
                "Google Drive вернул пустой файл."
            )

        target = unique_path(folder / result_path.name)
        shutil.move(str(result_path), str(target))

        log_result(
            log_path,
            item,
            "OK",
            target,
            target.stat().st_size,
            1,
            "Google Drive: файл скачан",
        )

        return [target]

    finally:
        shutil.rmtree(temp_folder, ignore_errors=True)
        
def download_youtube(url, folder, item, log_path):
    temp_folder = folder / f".youtube_tmp_{int(time.time() * 1000)}"
    temp_folder.mkdir(parents=True, exist_ok=True)

    output_template = str(temp_folder / "%(title)s.%(ext)s")
    proxy = "http://127.0.0.1:10809"

    def ps_quote(value):
        # Безопасное экранирование одинарных кавычек PowerShell.
        return "'" + str(value).replace("'", "''") + "'"

    ps_command = (
        f"& {ps_quote(sys.executable)} -m yt_dlp "
        f"--verbose "
        f"--proxy {ps_quote(proxy)} "
        f"--no-playlist "
        f"--retries 10 "
        f"--fragment-retries 10 "
        f"--file-access-retries 3 "
        f"--socket-timeout 30 "
        f"-f {ps_quote('best[ext=mp4]/best')} "
        f"-o {ps_quote(output_template)} "
        f"{ps_quote(url)}; "
        "exit $LASTEXITCODE"
    )

    last_error = ""

    try:
        for attempt in range(1, MAX_RETRIES + 1):
            print(f"\n[YouTube] Попытка {attempt}/{MAX_RETRIES}")
            print(f"[YouTube] Ссылка: {url}")

            try:
                # Очищаем остатки предыдущей неудачной попытки.
                for path in temp_folder.iterdir():
                    if path.is_file():
                        path.unlink()

                result = subprocess.run(
                    [
                        "powershell.exe",
                        "-NoProfile",
                        "-NonInteractive",
                        "-Command",
                        ps_command,
                    ],
                    capture_output=True,
                    text=True,
                    encoding="utf-8",
                    errors="replace",
                    timeout=600,
                    check=False,
                )

                if result.stdout:
                    print(result.stdout)

                if result.stderr:
                    print(result.stderr)

                if result.returncode != 0:
                    raise RuntimeError(
                        f"yt-dlp завершился с кодом {result.returncode}"
                    )

                downloaded = [
                    path for path in temp_folder.iterdir()
                    if path.is_file()
                    and not path.name.endswith(
                        (".part", ".ytdl", ".temp")
                    )
                ]

                if not downloaded:
                    raise RuntimeError(
                        "yt-dlp завершился без ошибки, "
                        "но итоговый файл не найден."
                    )

                saved_files = []

                for source in downloaded:
                    target = unique_path(folder / source.name)
                    shutil.move(str(source), str(target))
                    saved_files.append(target)

                total_size = sum(
                    path.stat().st_size for path in saved_files
                )

                log_result(
                    log_path,
                    item,
                    "OK",
                    saved_files[0],
                    total_size,
                    attempt,
                    f"YouTube: скачано файлов: {len(saved_files)}",
                )

                return saved_files

            except Exception as exc:
                last_error = f"{type(exc).__name__}: {exc}"
                print(f"[YouTube] Ошибка: {last_error}")

                if attempt < MAX_RETRIES:
                    print("[YouTube] Повтор через 3 секунды...")
                    time.sleep(3)

        log_result(
            log_path,
            item,
            "ERROR",
            attempts=MAX_RETRIES,
            details=f"YouTube: {last_error}",
        )

        return []

    finally:
        shutil.rmtree(temp_folder, ignore_errors=True)

    def is_youtube(url):
        host = (urlparse(url).hostname or "").lower()
        return (
            host == "youtu.be"
            or host == "youtube.com"
            or host.endswith(".youtube.com")
        )
# ============================================================
# Cloud@Mail.Ru
# ============================================================

MAILRU_DOWNLOADER_DIR = (
    Path(__file__).resolve().parent / "Cloud-Mail.Ru-Downloader"
)


def is_mailru(url):
    host = (urlparse(url).hostname or "").lower()
    return host in {
        "cloud.mail.ru",
        "www.cloud.mail.ru",
    }


def download_mailru(url, folder, item, log_path):
    """
    Запускает внешний Cloud-Mail.Ru-Downloader в отдельной
    временной копии. Готовые файлы переносит в папку материала.
    """
    import tempfile

    source_dir = MAILRU_DOWNLOADER_DIR

    if not source_dir.is_dir():
        raise FileNotFoundError(
            "Не найдена папка загрузчика Mail.Ru:\n"
            f"{source_dir}\n\n"
            "Положи папку Cloud-Mail.Ru-Downloader рядом "
            "с этим Python-скриптом."
        )

    required = [
        source_dir / "RUN.bat",
        source_dir / "links.txt",
        source_dir / "cloud_mail_downloader.php",
        source_dir / "aria2c.exe",
        source_dir / "php",
    ]

    missing = [str(p) for p in required if not p.exists()]
    if missing:
        raise FileNotFoundError(
            "В папке Cloud-Mail.Ru-Downloader не хватает файлов:\n"
            + "\n".join(missing)
        )

    folder.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory(prefix="dvizh_mailru_") as temp:
        work_dir = Path(temp) / "downloader"

        # Копируем сборку целиком, чтобы RUN.bat и PHP работали
        # со своими относительными путями, как при ручном запуске.
        shutil.copytree(
            source_dir,
            work_dir,
            ignore=shutil.ignore_patterns(
                "downloads",
                "input.txt",
                "*.aria2",
            ),
        )

        links_file = work_dir / "links.txt"
        links_file.write_text(
            url.strip() + "\n",
            encoding="utf-8",
        )

        print("[Mail.Ru] Запускаю внешний загрузчик...")
        print(f"[Mail.Ru] Ссылка: {url}")

        # Запускаем BAT из его собственной рабочей папки.
        # timeout намеренно не задан: большие видео могут
        # скачиваться долго.
        result = subprocess.run(
            ["cmd.exe", "/d", "/c", "RUN.bat"],
            cwd=str(work_dir),
            check=False,
        )

        download_dir = work_dir / "downloads"

        if result.returncode != 0:
            raise RuntimeError(
                "RUN.bat завершился с кодом "
                f"{result.returncode}."
            )

        if not download_dir.is_dir():
            raise RuntimeError(
                "Загрузчик завершился, но папка downloads "
                "не найдена."
            )

        # Забираем только готовые файлы, исключая служебные
        # и недокачанные файлы aria2.
        downloaded = [
            p for p in download_dir.rglob("*")
            if p.is_file()
            and p.stat().st_size > 0
            and not p.name.endswith((".aria2", ".part", ".tmp"))
        ]

        if not downloaded:
            raise RuntimeError(
                "Загрузчик не создал готовых файлов. "
                "Проверь вывод RUN.bat выше."
            )

        saved_files = []

        for source in downloaded:
            relative = source.relative_to(download_dir)
            destination_dir = folder / relative.parent
            destination_dir.mkdir(parents=True, exist_ok=True)

            target = unique_path(destination_dir / relative.name)
            shutil.move(str(source), str(target))
            saved_files.append(target)

        total_size = sum(p.stat().st_size for p in saved_files)

        log_result(
            log_path,
            item,
            "OK",
            saved_files[0],
            total_size,
            1,
            f"Mail.Ru: скачано файлов: {len(saved_files)}",
        )

        print(
            f"[Mail.Ru] Готово. Файлов: {len(saved_files)}, "
            f"размер: {total_size:,} байт"
        )

        return saved_files
        
def download_one(item, root_folder, session, log_path):
    category_folder = root_folder / safe_name(item["category"])
    request_folder = category_folder / safe_name(item["request_no"])
    material_folder = request_folder / safe_name(item["material"])
    material_folder.mkdir(parents=True, exist_ok=True)

    url = item["url"]

    try:
        if is_mailru(url):
            return download_mailru(
                url, material_folder, item, log_path
            )
        # YouTube проверяется отдельно до обычной HTTP-загрузки.
        if is_youtube(url):
            return download_youtube(
                url, material_folder, item, log_path
            )

        if is_google_drive(url):
            return download_google(
                url, material_folder, item, log_path
            )

        if is_yandex(url):
            direct_url, metadata = yandex_download_url(url, session)

            # Имя из метаданных Яндекс Диска обычно точнее подписи Excel.
            item_for_download = dict(item)
            if metadata.get("name"):
                item_for_download["material"] = metadata["name"]

            return download_http(
                direct_url, material_folder,
                session, item_for_download, log_path
            )

        # В том числе прямые ссылки на файлы Cosplay2.
        return download_http(
            url, material_folder, session, item, log_path
        )

    except Exception as exc:
        details = f"{type(exc).__name__}: {exc}"
        log_result(
            log_path, item, "ERROR",
            attempts=1, details=details
        )
        print("  Ошибка источника:", details)
        return []


def main():
    print("=" * 64)
    print(APP_NAME)
    print("=" * 64)

    print("\nВыбери Excel-таблицу со ссылками.")
    excel_path = choose_file()

    if not excel_path:
        print("Файл не выбран.")
        return

    print("\nЧитаю таблицу...")
    records = read_excel(excel_path)

    if not records:
        print("Не найдены строки со ссылками.")
        print(
            "Проверь, что есть столбец «Ссылка», «URL» или «Link» "
            "и ссылки записаны в ячейках или гиперссылках."
        )
        return

    print(f"Найдено уникальных ссылок: {len(records)}")
    print("\nВыбери папку назначения.")
    root_folder = choose_folder()

    if not root_folder:
        print("Папка назначения не выбрана.")
        return

    log_path = root_folder / "dvizh_download_log.csv"

    session = requests.Session()
    session.headers.update({"User-Agent": USER_AGENT})

    successes = 0
    failures = 0

    try:
        for index, item in enumerate(records, start=1):
            print("\n" + "-" * 64)
            print(f"[{index}/{len(records)}]")
            print(
                f"{item['category']} / "
                f"заявка {item['request_no']} / {item['material']}"
            )
            print(item["url"])

            files = download_one(
                item, root_folder, session, log_path
            )

            if files:
                successes += 1
            else:
                failures += 1

    finally:
        session.close()

    print("\n" + "=" * 64)
    print("ЗАГРУЗКА ЗАВЕРШЕНА")
    print(f"Ссылок обработано: {len(records)}")
    print(f"Успешно: {successes}")
    print(f"Ошибок: {failures}")
    print(f"Папка материалов: {root_folder}")
    print(f"Журнал: {log_path}")

    try:
        ui = tk.Tk()
        ui.withdraw()
        ui.attributes("-topmost", True)
        messagebox.showinfo(
            "DVIZH — загрузка завершена",
            f"Ссылок: {len(records)}\n"
            f"Успешно: {successes}\n"
            f"Ошибок: {failures}\n\n"
            f"Журнал:\n{log_path}"
        )
        ui.destroy()
    except Exception:
        pass


if __name__ == "__main__":
    try:
        main()
    except Exception:
        print("\nКРИТИЧЕСКАЯ ОШИБКА:")
        traceback.print_exc()

    input("\nНажми Enter для закрытия...")
